"""
RADAR DE VAGAS — curriculo.py (NOVO)

Para cada vaga, gera:
  • um CURRÍCULO ADAPTADO (arquivo .docx, editável, de coluna única, amigável para sistemas de triagem/ATS)
  • um KIT DE CANDIDATURA no Telegram: resumo profissional, palavras-chave da vaga,
    respostas prontas para perguntas comuns e as lacunas que a vaga aponta.

De onde vêm os fatos: fase3_perfil/perfil.json + texto do Curriculo.pdf (se existir na raiz).
O Gemini só pode usar o que estiver nesses dados. Depois da geração, o código CONFERE:
  • empresa que não está nos seus dados  → a experiência é removida
  • número/data que não está nos seus dados → o item é removido (ou o período é apagado)
  • competência, formação ou certificação que não aparece nos dados → removida
Tudo o que for removido aparece como aviso no Telegram. Mesmo assim, REVISE antes de enviar.

Como é acionado:
  • automático: vagas com score >= CURRICULO_AUTO_SCORE (padrão 85), até CURRICULO_AUTO_MAX por rodada (padrão 2)
  • sob demanda: botão "📄 Gerar currículo + kit" na mensagem da vaga (processado na próxima rodada,
    até CURRICULO_MAX_POR_RODADA por rodada, padrão 3)

Dependências: python-docx (e pypdf, que já está no workflow).
"""

import os
import re
import io
import json
import time
import contextlib
import unicodedata
from pathlib import Path
from datetime import datetime

from paginas_vaga import obter_descricao
from telegram_acoes import (enviar_telegram, enviar_documento, _esc, _planilha,
                            PENDENTE_CURRICULO, PRIMEIRA_LINHA_DADOS)

PASTA_SAIDA = Path(os.environ.get("CURRICULO_PASTA", "/tmp/curriculos"))
AUTO_SCORE = int(os.environ.get("CURRICULO_AUTO_SCORE", "85"))
AUTO_MAX = int(os.environ.get("CURRICULO_AUTO_MAX", "2"))
MAX_PENDENTES = int(os.environ.get("CURRICULO_MAX_POR_RODADA", "3"))
MODELOS = ['gemini-2.5-flash', 'gemini-3.5-flash']


# ====================================================================
# ── DADOS DO CANDIDATO (a única fonte de fatos) ──────────────────────
# ====================================================================

def _texto_curriculo_pdf() -> str:
    caminho = Path(os.environ.get("CURRICULO_PDF") or (Path(__file__).parent / "Curriculo.pdf"))
    if not caminho.exists():
        return ""
    try:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(caminho)).pages).strip()
    except Exception as e:
        print(f"  ⚠️  Não consegui ler o Curriculo.pdf: {e}")
        return ""


def _fonte_dados() -> str:
    perfil = {}
    try:
        from perfil import carregar_perfil  # fase3_perfil/perfil.py (o main.py já põe a pasta no caminho)
        with contextlib.redirect_stdout(io.StringIO()):
            perfil = carregar_perfil()
    except Exception as e:
        print(f"  ⚠️  Não consegui carregar o perfil.json: {e}")
    partes = []
    if perfil:
        partes.append("PERFIL (JSON):\n" + json.dumps(perfil, ensure_ascii=False))
    pdf = _texto_curriculo_pdf()
    if pdf:
        partes.append("CURRÍCULO ATUAL (texto do PDF):\n" + pdf)
    return "\n\n".join(partes)[:14000]


# ====================================================================
# ── GEMINI ───────────────────────────────────────────────────────────
# ====================================================================

def _gerar_json(vaga: dict, descricao: str, fonte: str):
    api_key = os.environ.get("GEMINI_API_KEY")
    if not api_key:
        print("  ⚠️  GEMINI_API_KEY não encontrada.")
        return None
    from google import genai
    client = genai.Client(api_key=api_key)

    prompt = f"""Você é especialista em currículos para sistemas de triagem (ATS) no Brasil.
Monte o currículo e o kit de candidatura do candidato para a vaga abaixo.

REGRAS DE OURO (não quebre nenhuma):
- Use SOMENTE fatos presentes em DADOS DO CANDIDATO. Não invente empregadores, cargos, datas, formações, ferramentas, números ou resultados.
- Você pode reordenar, reescrever e dar ênfase, e usar as palavras da vaga quando o candidato realmente tem aquela experiência.
- Se a vaga pede algo que o candidato não comprova, NÃO coloque no currículo: liste em "lacunas".
- Português do Brasil, direto, verbos de ação, sem exageros. Datas e nomes exatamente como estão nos dados.

VAGA
Título: {vaga.get('titulo', '')}
Empresa: {vaga.get('empresa', '')}
Local: {vaga.get('local', '')}
Descrição:
{descricao if descricao else '(indisponível: use apenas o título e deixe palavras_chave_vaga e lacunas vazias se não houver base)'}

DADOS DO CANDIDATO
{fonte}

Responda APENAS com um objeto JSON válido (começando com {{ e terminando com }}):
{{
  "nome": "nome completo, como está nos dados",
  "contatos": ["e-mail, telefone, cidade, LinkedIn, GitHub: só o que existir nos dados"],
  "titulo_profissional": "linha de posicionamento curta e honesta para esta vaga",
  "resumo_profissional": "3 a 4 linhas",
  "competencias": ["até 12 competências/ferramentas que o candidato TEM e que a vaga valoriza"],
  "experiencias": [{{"cargo": "", "empresa": "", "periodo": "", "bullets": ["até 5, os mais relevantes para a vaga primeiro"]}}],
  "formacao": ["Curso — Instituição (situação/ano)"],
  "certificacoes": ["..."],
  "palavras_chave_vaga": ["até 10 termos importantes da vaga"],
  "lacunas": ["requisitos da vaga que o candidato não comprova"],
  "respostas": {{
    "por_que_empresa": "2 a 3 frases",
    "por_que_voce": "2 a 3 frases",
    "experiencia_relevante": "2 a 3 frases"
  }}
}}"""

    for modelo in MODELOS:
        for tentativa in range(2):
            try:
                resp = client.models.generate_content(model=modelo, contents=prompt)
                achado = re.search(r'\{.*\}', resp.text.strip(), re.DOTALL)
                if not achado:
                    raise ValueError("JSON não encontrado")
                return json.loads(achado.group(0))
            except Exception as e:
                erro = str(e)
                if '429' in erro and tentativa == 0:
                    print(f"  ⏳ {modelo} — limite de requisições, aguardando 20s...")
                    time.sleep(20)
                    continue
                print(f"  ⚠️  {modelo} falhou: {erro[:80]}")
                break
    return None


# ====================================================================
# ── TRAVAS CONTRA INFORMAÇÃO INVENTADA ───────────────────────────────
# ====================================================================

def _norm(txt: str) -> str:
    t = unicodedata.normalize("NFD", (txt or "").lower())
    return "".join(c for c in t if unicodedata.category(c) != "Mn")


def _existe(texto: str, fonte_norm: str, minimo: float = 0.6) -> bool:
    """O texto aparece nos seus dados? (inteiro, ou pelo menos 60% das palavras com 4+ letras)."""
    n = _norm(texto).strip()
    if not n:
        return False
    if n in fonte_norm:
        return True
    tokens = [w for w in re.findall(r"\w+", n) if len(w) >= 4]
    if not tokens:
        return False
    return sum(1 for w in tokens if w in fonte_norm) / len(tokens) >= minimo


def _numeros_fora(texto: str, fonte: str) -> list:
    nums = [n.rstrip(".,%") for n in re.findall(r"\d[\d.,]*%?", texto or "")]
    return [n for n in nums if n and n not in fonte]


def _validar(dados: dict, fonte: str):
    fn = _norm(fonte)
    avisos = []

    comp = []
    for c in dados.get("competencias") or []:
        if isinstance(c, str) and _existe(c, fn):
            comp.append(c)
        else:
            avisos.append(f"Competência removida (não está nos seus dados): {c}")
    dados["competencias"] = comp

    exps = []
    for e in dados.get("experiencias") or []:
        if not isinstance(e, dict):
            continue
        empresa = (e.get("empresa") or "").strip()
        if empresa and _norm(empresa) not in fn:
            avisos.append(f"Experiência removida (empresa não está nos seus dados): {empresa}")
            continue
        fora = _numeros_fora(e.get("periodo", ""), fonte)
        if fora:
            avisos.append(f"Período apagado em '{empresa}' (datas não conferem com seus dados): confira")
            e["periodo"] = ""
        bullets = []
        for b in e.get("bullets") or []:
            fora = _numeros_fora(b, fonte)
            if fora:
                avisos.append(f"Item removido por conter número que não está nos seus dados ({', '.join(fora)})")
                continue
            bullets.append(b)
        e["bullets"] = bullets
        exps.append(e)
    dados["experiencias"] = exps

    for chave, rotulo in (("formacao", "Formação"), ("certificacoes", "Certificação")):
        ok = []
        for item in dados.get(chave) or []:
            if isinstance(item, str) and _existe(item, fn):
                ok.append(item)
            else:
                avisos.append(f"{rotulo} removida (não está nos seus dados): {item}")
        dados[chave] = ok

    def sem_numeros_inventados(texto: str, onde: str) -> str:
        """Tira as frases que trazem número/ano que não está nos seus dados."""
        frases = re.split(r"(?<=[.!?])\s+", (texto or "").strip())
        boas = [f for f in frases if f and not _numeros_fora(f, fonte)]
        if len(boas) < len([f for f in frases if f]):
            avisos.append(f"Frase removida de '{onde}' por conter número que não está nos seus dados")
        return " ".join(boas)

    dados["resumo_profissional"] = sem_numeros_inventados(dados.get("resumo_profissional", ""), "resumo profissional")
    resp = dados.get("respostas") or {}
    dados["respostas"] = {k: sem_numeros_inventados(v, k) for k, v in resp.items() if isinstance(v, str)}
    return dados, avisos


# ====================================================================
# ── ARQUIVO .DOCX ────────────────────────────────────────────────────
# ====================================================================

def _slug(txt: str, n: int = 30) -> str:
    t = unicodedata.normalize("NFKD", txt or "").encode("ascii", "ignore").decode()
    t = re.sub(r"[^A-Za-z0-9]+", "_", t).strip("_")
    return t[:n] or "vaga"


def _salvar_docx(dados: dict, vaga: dict) -> str:
    from docx import Document
    from docx.shared import Pt, Cm

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    def paragrafo(texto="", negrito=False, tamanho=None, antes=0, depois=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(depois)
        if texto:
            r = p.add_run(texto)
            r.bold = negrito
            if tamanho:
                r.font.size = Pt(tamanho)
        return p

    def secao(titulo):
        paragrafo(titulo.upper(), negrito=True, tamanho=11, antes=8, depois=2)

    paragrafo(dados.get("nome") or "", negrito=True, tamanho=18, depois=0)
    if dados.get("titulo_profissional"):
        paragrafo(dados["titulo_profissional"], tamanho=11, depois=0)
    contatos = [c for c in dados.get("contatos") or [] if isinstance(c, str) and c.strip()]
    if contatos:
        paragrafo(" | ".join(contatos), depois=4)

    if dados.get("resumo_profissional"):
        secao("Resumo profissional")
        paragrafo(dados["resumo_profissional"])
    if dados.get("competencias"):
        secao("Competências")
        paragrafo(" • ".join(dados["competencias"]))
    if dados.get("experiencias"):
        secao("Experiência profissional")
        for e in dados["experiencias"]:
            cab = " — ".join(x for x in [e.get("cargo", ""), e.get("empresa", "")] if x)
            if e.get("periodo"):
                cab += f"  |  {e['periodo']}"
            paragrafo(cab, negrito=True, antes=4, depois=1)
            for b in e.get("bullets") or []:
                doc.add_paragraph(b, style="List Bullet").paragraph_format.space_after = Pt(0)
    if dados.get("formacao"):
        secao("Formação")
        for f in dados["formacao"]:
            paragrafo(f, depois=1)
    if dados.get("certificacoes"):
        secao("Certificações e cursos")
        for c in dados["certificacoes"]:
            doc.add_paragraph(c, style="List Bullet").paragraph_format.space_after = Pt(0)

    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)
    nome = f"Curriculo_{_slug(dados.get('nome', ''), 20)}_{_slug(vaga.get('empresa', ''), 20)}_{_slug(vaga.get('titulo', ''), 30)}.docx"
    caminho = str(PASTA_SAIDA / nome)
    doc.save(caminho)
    return caminho


# ====================================================================
# ── KIT NO TELEGRAM ──────────────────────────────────────────────────
# ====================================================================

def _mensagens_kit(dados: dict, vaga: dict, avisos: list) -> list[str]:
    r = dados.get("respostas") or {}
    t = f"🧰 <b>KIT DE CANDIDATURA</b>\n<b>{_esc(vaga.get('titulo', ''))}</b> — {_esc(vaga.get('empresa', ''))}\n\n"
    if dados.get("resumo_profissional"):
        t += f"📝 <b>Resumo profissional</b>\n{_esc(dados['resumo_profissional'])[:700]}\n\n"
    if dados.get("palavras_chave_vaga"):
        t += "🔑 <b>Palavras-chave da vaga</b>\n" + ", ".join(_esc(p) for p in dados["palavras_chave_vaga"][:10]) + "\n\n"
    p2 = ""
    for chave, rotulo in (("por_que_empresa", "Por que esta empresa?"),
                          ("por_que_voce", "Por que você?"),
                          ("experiencia_relevante", "Experiência mais relevante")):
        if r.get(chave):
            p2 += f"❓ <b>{rotulo}</b>\n{_esc(r[chave])[:600]}\n\n"
    if dados.get("lacunas"):
        p2 += "🔸 <b>Lacunas da vaga (não incluídas no currículo)</b>\n" + "; ".join(_esc(x)[:150] for x in dados["lacunas"][:6]) + "\n\n"
    if avisos:
        p2 += "⚠️ <b>Conferir</b>\n" + "\n".join(f"• {_esc(a)[:200]}" for a in avisos[:8]) + "\n\n"
    p2 += "<i>Revise antes de enviar: confira datas, empresas e números. Nada foi enviado por você.</i>"
    if len(t) + len(p2) <= 3900:
        return [t + p2]
    return [t.strip(), p2]


# ====================================================================
# ── FLUXO PRINCIPAL ──────────────────────────────────────────────────
# ====================================================================

def gerar_e_enviar(vaga: dict, token: str, chat_id: str) -> bool:
    """Gera o currículo adaptado + kit e envia ao Telegram. True se enviou."""
    print(f"  📄 Gerando currículo adaptado: '{(vaga.get('titulo') or '')[:40]}'...")
    fonte = _fonte_dados()
    if not fonte.strip():
        print("  ⚠️  Sem dados do candidato (perfil.json / Curriculo.pdf). Currículo não gerado.")
        return False
    descricao = (vaga.get("descricao") or "").strip() or obter_descricao(vaga)
    dados = _gerar_json(vaga, descricao, fonte)
    if not dados:
        print("  ❌ O Gemini não devolveu o currículo. Tenta de novo na próxima rodada.")
        return False

    dados, avisos = _validar(dados, fonte)
    try:
        caminho = _salvar_docx(dados, vaga)
    except Exception as e:
        print(f"  ❌ Não consegui montar o .docx ({type(e).__name__}): {e}")
        return False

    legenda = (f"📄 <b>Currículo adaptado</b>\n{_esc(vaga.get('titulo', ''))[:150]} — "
               f"{_esc(vaga.get('empresa', ''))}\n<i>Revise antes de enviar.</i>")
    if not enviar_documento(caminho, legenda, token, chat_id):
        return False
    for msg in _mensagens_kit(dados, vaga, avisos):
        enviar_telegram(msg, token, chat_id)
    print(f"  ✅ Currículo + kit enviados ({len(avisos)} aviso(s) de conferência)")
    return True


def gerar_curriculos_pendentes(token: str, chat_id: str) -> None:
    """Atende os cliques em '📄 Gerar currículo + kit' (a coluna L da planilha guarda o pedido)."""
    print("\n[CURRÍCULO] Procurando pedidos de currículo na planilha...")
    aba = _planilha()
    if aba is None:
        return
    pendentes = []
    for n, lin in enumerate(aba.get_all_values(), start=1):
        if n < PRIMEIRA_LINHA_DADOS:
            continue
        lin = lin + [""] * (15 - len(lin))
        if lin[11].strip() == PENDENTE_CURRICULO and lin[2].strip():
            pendentes.append((n, lin))
        if len(pendentes) >= MAX_PENDENTES:
            break
    if not pendentes:
        print("  🤫 Nenhum pedido pendente.")
        return

    hoje = datetime.now().strftime("%d/%m")
    for n, lin in pendentes:
        vaga = {"titulo": lin[1], "empresa": lin[3], "link": lin[2], "fonte": lin[7], "local": lin[6]}
        if gerar_e_enviar(vaga, token, chat_id):
            try:
                aba.update(range_name=f"L{n}", values=[[f"📄 Currículo enviado em {hoje}"]],
                           value_input_option="USER_ENTERED")
            except Exception as e:
                print(f"  ⚠️  Enviei, mas não consegui marcar a planilha: {e}")
